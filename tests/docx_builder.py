import base64
from io import BytesIO

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Inches

TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGA"
    "hKmMIQAAAABJRU5ErkJggg=="
)


def make_docx(lines):
    document = Document()
    for line in lines:
        document.add_paragraph(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_with_table(rows):
    document = Document()
    document.add_paragraph("Goutham Sai")
    document.add_paragraph("goutham@example.com | +1 555 123 4567")

    table = document.add_table(rows=0, cols=len(rows[0]))
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_with_header_contact(text):
    document = Document()
    document.sections[0].header.paragraphs[0].text = text
    document.add_paragraph("Goutham Sai")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_with_header_table_contact(email):
    document = Document()
    header = document.sections[0].header
    table = header.add_table(rows=1, cols=2, width=Inches(6))
    table.cell(0, 0).text = "Goutham Sai"
    table.cell(0, 1).text = email

    document.add_paragraph("Goutham Sai")
    document.add_paragraph("Experience")
    document.add_paragraph("Built REST APIs in Python for reporting tools.")
    document.add_paragraph("Education")
    document.add_paragraph("BSc Computer Science")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_with_image(width_inches=0.4, height_inches=0.4, body_lines=None):
    document = Document()
    for line in body_lines or []:
        document.add_paragraph(line)

    document.add_picture(
        BytesIO(TINY_PNG), width=Inches(width_inches), height=Inches(height_inches)
    )

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_with_textbox(box_text, body_lines=None):
    document = Document()
    for line in body_lines or []:
        document.add_paragraph(line)

    paragraph = document.add_paragraph("Lead paragraph for the document.")
    run = paragraph.add_run()
    xml = (
        '<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        '2006/main" xmlns:v="urn:schemas-microsoft-com:vml">'
        '<v:shape id="TextBox1" style="width:200pt;height:30pt">'
        "<v:textbox><w:txbxContent>"
        f"<w:p><w:r><w:t>{box_text}</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox>"
        "</v:shape>"
        "</w:pict>"
    )
    run._r.append(parse_xml(xml))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
