from app.services.extractor import extract_skills
from app.services.parser import (
    extract_text_from_docx_bytes,
    extract_text_from_pdf_bytes,
    normalize_extracted_text,
)
from tests.docx_builder import (
    make_docx,
    make_docx_with_header_contact,
    make_docx_with_table,
    make_docx_with_textbox,
)
from tests.pdf_builder import make_pdf


def test_extract_text_from_invalid_pdf_bytes_returns_empty():
    bad_bytes = b"this is not a real pdf"

    assert extract_text_from_pdf_bytes(bad_bytes) == ""


def test_extract_text_from_truncated_pdf_returns_empty():
    full = make_pdf([(50, 60, "Hello resume")])
    truncated = full[:200]

    assert extract_text_from_pdf_bytes(truncated) == ""


def test_extract_text_from_valid_pdf():
    data = make_pdf([(50, 60, "Python developer with FastAPI experience")])
    text = extract_text_from_pdf_bytes(data)

    assert "Python developer with FastAPI experience" in text


def test_extract_text_fixes_colons_glued_by_pdf_extraction():
    data = make_pdf([(50, 60, "Marketing and Community:Publicity Campaigns")])
    text = extract_text_from_pdf_bytes(data)

    assert "Community: Publicity" in text


def test_extract_text_from_invalid_docx_bytes_returns_empty():
    assert extract_text_from_docx_bytes(b"this is not a real docx") == ""


def test_extract_text_from_docx_includes_paragraphs():
    data = make_docx(["Goutham Sai", "Experience", "Built REST APIs in Python."])
    text = extract_text_from_docx_bytes(data)

    assert "Goutham Sai" in text
    assert "Built REST APIs in Python." in text


def test_extract_text_from_docx_includes_table_cells():
    data = make_docx_with_table([["Python", "5 years"], ["FastAPI", "3 years"]])
    text = extract_text_from_docx_bytes(data)

    assert "Python | 5 years" in text
    assert "FastAPI | 3 years" in text


def test_extract_text_repairs_period_glue_from_pdf():
    data = make_pdf([(50, 60, "Built services.Python apps with FastAPI.")])
    text = extract_text_from_pdf_bytes(data)

    assert "services. Python" in text


def test_glued_period_text_does_not_lose_skills():
    repaired = normalize_extracted_text(
        "Built services.Python apps with FastAPI and Next.js."
    )

    skills = extract_skills(repaired)
    assert "python" in skills
    assert "fastapi" in skills


def test_text_normalization_keeps_versions_urls_and_domains_intact():
    text = normalize_extracted_text(
        "Python 3.12, node.js, and https://example.com/docs stay intact."
    )

    assert "3.12" in text
    assert "node.js" in text
    assert "https://example.com/docs" in text


def test_pdf_multi_column_pages_read_rows_in_order():
    items = []
    for index in range(14):
        top = 100 + index * 18
        items.append((50, top, f"left row {index} built scalable backend systems"))
        items.append(
            (
                350,
                top,
                f"right row {index} python fastapi docker kubernetes observability",
            )
        )

    text = extract_text_from_pdf_bytes(make_pdf(items))

    assert "left row 13" in text
    assert text.index("left row 0") < text.index("right row 0")
    assert text.index("right row 0") < text.index("left row 1")
    assert text.index("left row 13") < text.index("right row 13")


def test_docx_text_boxes_extracted_correctly():
    data = make_docx_with_textbox(
        "InsideTextBoxVisibleText", ["Goutham Sai", "Python developer"]
    )

    text = extract_text_from_docx_bytes(data)

    assert "InsideTextBoxVisibleText" in text
    assert "Lead paragraph for the document." in text
    assert "Goutham Sai" in text


def test_tight_kerning_pdf_deglued_properly():
    items = [
        (50, 60, "Goutham Sai"),
        (50, 100, "Kept detailed records"),
        (139.7, 100, "for every project milestone."),
    ]

    text = extract_text_from_pdf_bytes(make_pdf(items))

    assert "records for every project milestone" in text
    assert "recordsfor" not in text


def test_glued_words_split_only_when_both_halves_are_real_words():
    repaired = normalize_extracted_text(
        "Kept recordsfor the team and met managementskills goals.",
        deglue=True,
    )

    assert "records for" in repaired
    assert "management skills" in repaired
    assert "gouthamsai" in normalize_extracted_text("gouthamsai", deglue=True)


def test_invisible_characters_are_stripped():
    text = normalize_extracted_text("Pyth\u200b\u200bon and Fast\u00adAPI")

    assert text == "Python and FastAPI"


def test_docx_header_text_is_extracted():
    data = make_docx_with_header_contact("goutham@example.com | +1 555 123 4567")

    text = extract_text_from_docx_bytes(data)

    assert "goutham@example.com" in text


def test_docx_tables_keep_document_order():
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_paragraph("Alpha body line for testing order.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Beta table cell"
    table.rows[0].cells[1].text = "Gamma table cell"
    document.add_paragraph("Delta body line for testing order.")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_text_from_docx_bytes(buffer.getvalue())

    assert text.index("Alpha") < text.index("Beta") < text.index("Delta")
