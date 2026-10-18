from io import BytesIO

from docx import Document

from app.services.tailor import bridge_aliases, build_resume_docx


def test_bridge_aliases_replaces_with_canonical_names():
    result = bridge_aliases(
        "Worked with k8s clusters and postgres databases daily.",
        ["kubernetes", "postgresql"],
    )

    assert "Kubernetes (k8s)" in result["text"]
    assert "PostgreSQL (postgres)" in result["text"]
    assert len(result["changes"]) == 2


def test_bridge_aliases_ignores_units_and_case_mismatches():
    result = bridge_aliases(
        "Collected 5 ml samples using 3 dl flasks. This ml pipeline handles data.",
        ["machine learning", "deep learning"],
    )

    assert "5 ml samples" in result["text"]
    assert "3 dl flasks" in result["text"]
    assert "This ml pipeline" in result["text"]
    assert result["changes"] == []

    bridged = bridge_aliases(
        "Our ML models and DL pipelines scale well.",
        ["machine learning", "deep learning"],
    )

    assert "Machine Learning (ML)" in bridged["text"]
    assert "Deep Learning (DL)" in bridged["text"]
    assert len(bridged["changes"]) == 2


def test_bridge_aliases_leaves_exact_matches_alone():
    result = bridge_aliases("Deployed services on Kubernetes.", ["kubernetes"])

    assert result["text"] == "Deployed services on Kubernetes."
    assert result["changes"] == []


def test_build_resume_docx_creates_structured_document():
    resume = (
        "Goutham Sai\n"
        "goutham@example.com | +1 555 123 4567\n"
        "Experience\n"
        "- Built REST APIs in Python for three years.\n"
        "- Led migration of 12 services to Docker.\n"
        "Skills\n"
        "Python, FastAPI, SQL, Docker\n"
        "Education\n"
        "BSc Computer Science, 2022\n"
    )

    data = build_resume_docx(resume)

    assert data[:2] == b"PK"

    document = Document(BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)

    assert "Goutham Sai" in text
    assert "EXPERIENCE" in text
    assert "Built REST APIs in Python for three years." in text
    assert "SKILLS" in text
    assert "EDUCATION" in text

    bullet_styles = [
        p.style.name for p in document.paragraphs if "Built REST" in p.text
    ]
    assert bullet_styles == ["List Bullet"]


def test_build_resume_docx_handles_plain_text():
    data = build_resume_docx("just some text without any sections at all")
    document = Document(BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)

    assert "just some text without any sections at all" in text


def test_build_resume_docx_does_not_treat_contact_line_as_name():
    resume = (
        "goutham@example.com | +1 555 123 4567\n"
        "Experience\n"
        "- Built REST APIs in Python for three years.\n"
    )

    data = build_resume_docx(resume)
    document = Document(BytesIO(data))
    first = document.paragraphs[0]

    assert "goutham@example.com" in first.text
    assert not any(run.bold for run in first.runs)


def test_build_resume_docx_keeps_summary_out_of_contact_line():
    resume = (
        "Jordan Lee\n"
        "Senior Marketing Manager\n"
        "jordan.lee@example.com | +1 555 123 4567\n"
        "Results-driven marketer with 8 years of experience growing pipelines.\n"
        "Experience\n"
        "- Grew organic traffic by 45 percent.\n"
    )

    data = build_resume_docx(resume)
    document = Document(BytesIO(data))
    paragraphs = document.paragraphs

    name_paragraph = paragraphs[0]
    assert name_paragraph.text == "Jordan Lee"
    assert any(run.bold for run in name_paragraph.runs)

    headline_paragraph = paragraphs[1]
    assert headline_paragraph.text == "Senior Marketing Manager"
    assert not any(run.bold for run in headline_paragraph.runs)

    contact_paragraph = paragraphs[2]
    assert "jordan.lee@example.com" in contact_paragraph.text
    assert "Results-driven" not in contact_paragraph.text

    assert any(p.text == "SUMMARY" for p in paragraphs)
    assert any(p.text.startswith("Results-driven marketer") for p in paragraphs)


def test_build_resume_docx_normalizes_skills_and_tools_heading():
    resume = (
        "Goutham Sai\n"
        "Skills & Tools\n"
        "Python, FastAPI, SQL\n"
        "Education\n"
        "BSc Computer Science, 2022\n"
    )

    data = build_resume_docx(resume)
    document = Document(BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)

    assert "SKILLS" in text
    assert "Python, FastAPI, SQL" in text
