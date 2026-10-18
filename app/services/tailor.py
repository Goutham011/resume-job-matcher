import io
import re

from docx import Document
from docx.shared import Inches, Pt

from app.services.contacts import EMAIL_RE, LINK_RE, looks_like_phone
from app.services.headings import HEADING_GROUPS, match_heading
from app.services.keywords import (
    CASE_SENSITIVE_SPELLINGS,
    SKILL_ALIASES,
    display_name,
    term_occurs,
)

BULLET_PREFIX_RE = re.compile(r"^[\s•\-\*\u2022\u25cf\u25cb\u25aa\u25e6]+")
NAME_LABEL_RE = re.compile(
    r"(email|phone|address|linkedin|github|portfolio)", re.IGNORECASE
)
TITLE_WORD_RE = re.compile(
    r"\b(?:manager|engineer|developer|designer|analyst|consultant|intern|"
    r"specialist|director|lead|architect|scientist|nurse|teacher|accountant|"
    r"marketer|administrator|coordinator|assistant|associate|executive|officer|"
    r"representative|technician|supervisor|strategist|writer|editor|producer|"
    r"planner|researcher|recruiter|generalist|enthusiast|professional)\b",
    re.IGNORECASE,
)
HEADLINE_MAX_WORDS = 8
CONTACT_LABEL_RE = re.compile(
    r"(email|phone|address|linkedin|github|portfolio)", re.IGNORECASE
)


def _looks_like_name(line):
    stripped = line.strip()
    if not stripped or len(stripped) > 60:
        return False
    if EMAIL_RE.search(stripped) or any(ch in stripped for ch in "|@/"):
        return False
    if any(ch.isdigit() for ch in stripped):
        return False
    if NAME_LABEL_RE.search(stripped):
        return False
    if TITLE_WORD_RE.search(stripped):
        return False
    if match_heading(stripped) in HEADING_GROUPS:
        return False
    words = stripped.split()
    if not 1 <= len(words) <= 5:
        return False
    return bool(re.search(r"[A-Za-zÀ-ÿ]", stripped))


def _is_contact_line(line):
    if EMAIL_RE.search(line) or LINK_RE.search(line) or looks_like_phone(line):
        return True
    return bool(CONTACT_LABEL_RE.search(line))


def _looks_like_headline(line):
    stripped = line.strip()
    if not stripped or len(stripped) > 60:
        return False
    if stripped.endswith((".", "!", "?")):
        return False
    if EMAIL_RE.search(stripped) or any(ch.isdigit() for ch in stripped):
        return False
    if "," in stripped:
        return False
    words = stripped.split()
    return 1 <= len(words) <= HEADLINE_MAX_WORDS


def _parse_resume(lines):
    name = None
    contact = []
    front = []
    sections = []
    current = None

    for raw in lines:
        line = raw.strip()
        if not line:
            continue

        title = match_heading(line)
        if title:
            current = {"title": title, "lines": []}
            sections.append(current)
            continue

        is_bullet = bool(BULLET_PREFIX_RE.match(line))
        text = BULLET_PREFIX_RE.sub("", line).strip()
        if current is not None:
            current["lines"].append((text, is_bullet))
        elif _is_contact_line(text):
            contact.append(text)
        else:
            front.append(text)

    name_index = None
    for index, text in enumerate(front):
        if _looks_like_name(text):
            name = text
            name_index = index
            break

    headline = None
    summary = []
    for index, text in enumerate(front):
        if index == name_index:
            continue
        adjacent = name_index is not None and abs(index - name_index) == 1
        if headline is None and adjacent and _looks_like_headline(text):
            headline = text
            continue
        summary.append(text)

    if summary:
        summary_lines = [(text, False) for text in summary]
        for section in sections:
            if section["title"] == "Summary":
                section["lines"][:0] = summary_lines
                break
        else:
            sections.insert(0, {"title": "Summary", "lines": summary_lines})

    return name, headline, contact, sections


def _alias_pattern(alias):
    spellings = CASE_SENSITIVE_SPELLINGS.get(alias)
    if spellings:
        alternation = "|".join(re.escape(spelling) for spelling in spellings)
        return re.compile(rf"(?<![\w.])(?:{alternation})(?![\w])")

    if len(alias) <= 2:
        return re.compile(rf"(?<![\w.]){re.escape(alias.upper())}(?![\w])")

    return re.compile(rf"(?<![\w.])(?<!\d ){re.escape(alias)}(?![\w])", re.IGNORECASE)


def bridge_aliases(resume_text, jd_skills):
    text = resume_text
    changes = []

    for skill in jd_skills:
        if term_occurs(skill, text):
            continue

        for alias in SKILL_ALIASES.get(skill, []):
            pattern = _alias_pattern(alias)

            def repl(match, target=skill):
                return f"{display_name(target)} ({match.group(0)})"

            new_text, replacements = pattern.subn(repl, text)
            if replacements:
                text = new_text
                changes.append(
                    {
                        "skill": skill,
                        "display": display_name(skill),
                        "alias": alias,
                        "replacements": replacements,
                    }
                )
                break

    return {"text": text, "changes": changes}


def build_resume_docx(resume_text):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    name, headline, contact, sections = _parse_resume(resume_text.splitlines())

    if name:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(name)
        run.bold = True
        run.font.size = Pt(16)

    if headline:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(headline)
        run.font.size = Pt(11)

    if contact:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run("  |  ".join(contact))
        run.font.size = Pt(10)

    for section in sections:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(section["title"].upper())
        run.bold = True
        run.font.size = Pt(11.5)

        for text, is_bullet in section["lines"]:
            if is_bullet:
                doc.add_paragraph(text, style="List Bullet")
            else:
                doc.add_paragraph(text)

    if not name and not sections and not contact:
        for line in resume_text.splitlines():
            line = line.strip()
            if line:
                doc.add_paragraph(line)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
